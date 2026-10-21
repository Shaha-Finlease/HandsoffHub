import os
import re
import json
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Load configuration
CONFIG_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'config.json')

def load_config():
    if os.path.exists(CONFIG_PATH):
        try:
            with open(CONFIG_PATH, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            pass
    return {
        "default_agency": "INH",
        "agency_manager_mapping": {},
        "portfolio_mapping": {}
    }

def save_config(cfg):
    with open(CONFIG_PATH, 'w', encoding='utf-8') as f:
        json.dump(cfg, f, indent=2)

def clean_str(val):
    if val is None:
        return ""
    s = str(val).strip()
    return "" if s.lower() == "none" else s

def to_float(val):
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).replace(',', '').strip()
    try:
        return float(s)
    except Exception:
        return 0.0

def format_date_str(val):
    if not val:
        return ""
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.strftime("%d-%m-%Y")
    s = str(val).strip()
    # Try multiple standard formats
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%b-%Y", "%d/%b/%Y",
                "%d-%m-%Y %H:%M:%S", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S",
                "%d-%b-%Y %H:%M:%S", "%d/%b/%Y %H:%M:%S %p", "%d/%m/%Y %I:%M:%S %p"):
        try:
            dt = datetime.datetime.strptime(s, fmt)
            return dt.strftime("%d-%m-%Y")
        except Exception:
            pass
    # If format contains date at start (e.g. 01/10/2026 06:58:32)
    m = re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{4})", s)
    if m:
        d, m, y = m.groups()
        return f"{int(d):02d}-{int(m):02d}-{y}"
    return s

def extract_cheque_no(narration, raw_cheque=None):
    if raw_cheque:
        s = clean_str(raw_cheque)
        if s and s not in ("0", "-", "None", "nan"):
            return s
            
    narration = clean_str(narration)
    if not narration:
        return ""

    # Common patterns:
    # 1. 12-digit UTR/RRN (Standard for UPI / IMPS)
    m12 = re.findall(r'\b\d{12}\b', narration)
    if m12:
        return m12[0]

    # 2. NEFT UTR e.g. SBIN526244349613, HDFCR52026..., AXISR...
    m_neft = re.search(r'\b[A-Z]{4}[0-9A-Z]{10,18}\b', narration)
    if m_neft:
        # Avoid bank names or standard words
        cand = m_neft.group(0)
        if not cand.startswith("SHAHA") and not cand.startswith("FINLEASE"):
            return cand

    # 3. Trailing cheque/deposit number e.g. CASH DEPOSIT ... 833
    m_cash = re.search(r'CASH DEPOSIT.*?[- ](\d{3,8})\b', narration, re.IGNORECASE)
    if m_cash:
        return m_cash.group(1)

    return ""

def detect_mode(mode_col, narration):
    m = clean_str(mode_col).upper()
    if m in ("UPI", "NEFT", "RTGS", "IMPS", "CASH", "CHEQUE", "CHQ", "ONLINE"):
        return m
    
    narr = clean_str(narration).upper()
    if "UPI" in narr:
        return "UPI"
    elif "NEFT" in narr:
        return "NEFT"
    elif "RTGS" in narr:
        return "RTGS"
    elif "IMPS" in narr:
        return "IMPS"
    elif "CASH" in narr:
        return "CASH"
    elif "CHEQUE" in narr or "CHQ" in narr:
        return "CHEQUE"
    elif "PAYMENT" in narr or "RAZORPAY" in narr:
        return "ONLINE"
    return "UPI" if not m else m

def parse_account_transactions(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    transactions = []
    
    for sheet_name in wb.sheetnames:
        # Skip summary sheets like 'Today' or sheets with no transactions
        s_lower = sheet_name.strip().lower()
        if s_lower in ("today", "summary", "dashboard"):
            continue
            
        bank_name = "HDFC" if "hdfc" in s_lower else ("ICICI" if "icici" in s_lower else ("Razorpay" if "razor" in s_lower else sheet_name))
        ws = wb[sheet_name]
        
        headers = []
        rows = ws.iter_rows(values_only=True)
        try:
            first_row = next(rows)
            headers = [clean_str(c).lower() for c in first_row if c is not None]
        except StopIteration:
            continue
            
        # Re-fetch row 1 as tuple
        wb_full_rows = list(ws.iter_rows(values_only=True))
        if not wb_full_rows:
            continue
            
        header_vals = [clean_str(c).lower() for c in wb_full_rows[0]]
        
        # Identify column indices
        col_ref = -1
        col_date = -1
        col_narration = -1
        col_amount = -1
        col_cheque = -1
        
        for idx, h in enumerate(header_vals):
            if not h: continue
            if "ref" in h and col_ref == -1:
                col_ref = idx
            elif any(d in h for d in ("date", "dt")) and col_date == -1:
                col_date = idx
            elif any(n in h for n in ("narration", "remark", "description")) and col_narration == -1:
                col_narration = idx
            elif any(a in h for a in ("credit", "amount", "tran. amount")) and col_amount == -1:
                col_amount = idx
            elif any(c in h for c in ("cheque", "chq", "tran. id", "utr")) and col_cheque == -1:
                col_cheque = idx
                
        # Fallbacks if columns not strictly matched
        if col_ref == -1: col_ref = 0
        if col_date == -1: col_date = 1
        if col_narration == -1: col_narration = 2
        if col_amount == -1: col_amount = 3
        
        for r_idx, r in enumerate(wb_full_rows[1:], start=2):
            if not any(r): continue
            
            ref_val = clean_str(r[col_ref]) if col_ref < len(r) else ""
            date_val = format_date_str(r[col_date]) if col_date < len(r) else ""
            narr_val = clean_str(r[col_narration]) if col_narration < len(r) else ""
            amt_val = to_float(r[col_amount]) if col_amount < len(r) else 0.0
            cheque_val = clean_str(r[col_cheque]) if (col_cheque != -1 and col_cheque < len(r)) else ""
            
            # If ref is empty, try generating or skip if empty row
            if not ref_val and amt_val <= 0 and not narr_val:
                continue
                
            transactions.append({
                "id": f"{sheet_name}_{r_idx}",
                "ref": ref_val,
                "date": date_val,
                "narration": narr_val,
                "amount": amt_val,
                "cheque_ref": cheque_val,
                "bank": bank_name,
                "sheet": sheet_name,
                "row_idx": r_idx,
                "allocated": False
            })
            
    wb.close()
    return transactions

def parse_paid_file(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    rows_data = []
    
    # Active sheet or sheet named 'Paid File'
    sheet_name = 'Paid File' if 'Paid File' in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    
    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        wb.close()
        return rows_data
        
    header_row = all_rows[0]
    headers = [clean_str(c).upper().replace('_', ' ') for c in header_row]
    
    # Standard column mapping
    # 'CUSTOMER ID ', 'CUSTOMER  NAME', 'PORTFOLIO', 'Date of Pay', 'Amount', 'MODE', 'TRANSACTION DESCRIPTION', 'BANK NAME', 'OFFICER NAME', 'MANAGER'
    idx_cust_id = -1
    idx_cust_name = -1
    idx_portfolio = -1
    idx_date = -1
    idx_amount = -1
    idx_mode = -1
    idx_narr = -1
    idx_bank = -1
    idx_officer = -1
    idx_manager = -1
    
    for idx, h in enumerate(headers):
        h_norm = " ".join(h.split())
        if any(k in h_norm for k in ("CUSTOMER ID", "LOAN ID", "LOAN NO")):
            idx_cust_id = idx
        elif any(k in h_norm for k in ("OFFICER", "CALLER")):
            idx_officer = idx
        elif "CUSTOMER NAME" in h_norm:
            idx_cust_name = idx
        elif "PORTFOLIO" in h_norm:
            idx_portfolio = idx
        elif any(k in h_norm for k in ("DATE OF PAY", "PAY DATE", "DATE")):
            idx_date = idx
        elif any(k in h_norm for k in ("AMOUNT", "AMT")):
            idx_amount = idx
        elif "MODE" in h_norm:
            idx_mode = idx
        elif any(k in h_norm for k in ("TRANSACTION DESCRIPTION", "DESCRIPTION", "NARRATION", "REMARKS")):
            idx_narr = idx
        elif "BANK" in h_norm:
            idx_bank = idx
        elif "MANAGER" in h_norm:
            idx_manager = idx
        elif "NAME" in h_norm and idx_cust_name == -1:
            idx_cust_name = idx
            
    for r_idx, r in enumerate(all_rows[1:], start=2):
        if not any(r): continue
        
        cust_id = clean_str(r[idx_cust_id]) if idx_cust_id != -1 and idx_cust_id < len(r) else ""
        cust_name = clean_str(r[idx_cust_name]) if idx_cust_name != -1 and idx_cust_name < len(r) else ""
        portfolio = clean_str(r[idx_portfolio]) if idx_portfolio != -1 and idx_portfolio < len(r) else ""
        date_of_pay = format_date_str(r[idx_date]) if idx_date != -1 and idx_date < len(r) else ""
        amt = to_float(r[idx_amount]) if idx_amount != -1 and idx_amount < len(r) else 0.0
        mode = clean_str(r[idx_mode]) if idx_mode != -1 and idx_mode < len(r) else ""
        narr = clean_str(r[idx_narr]) if idx_narr != -1 and idx_narr < len(r) else ""
        bank = clean_str(r[idx_bank]) if idx_bank != -1 and idx_bank < len(r) else ""
        officer = clean_str(r[idx_officer]) if idx_officer != -1 and idx_officer < len(r) else ""
        manager = clean_str(r[idx_manager]) if idx_manager != -1 and idx_manager < len(r) else ""
        
        if not cust_name and amt == 0.0 and not cust_id:
            continue
            
        rows_data.append({
            "customer_id": cust_id,
            "customer_name": cust_name,
            "portfolio": portfolio,
            "date_of_pay": date_of_pay,
            "amount": amt,
            "mode": mode,
            "narration": narr,
            "bank_name": bank,
            "officer_name": officer,
            "manager": manager,
            "source_file": os.path.basename(file_path),
            "row_num": r_idx
        })
        
    wb.close()
    return rows_data

def match_and_generate(account_transactions, paid_rows, handoff_date=None, config=None):
    if not config:
        config = load_config()
        
    if not handoff_date:
        handoff_date = datetime.datetime.now().strftime("%d-%m-%Y")
    else:
        handoff_date = format_date_str(handoff_date)
        
    portfolio_map = config.get("portfolio_mapping", {})
    manager_agency_map = config.get("agency_manager_mapping", {})
    default_agency = config.get("default_agency", "INH")
    
    # Tracking matched account transactions to avoid double-allocation
    used_tx_ids = set()
    handsoff_results = []
    
    for p_idx, p in enumerate(paid_rows):
        matched_tx = None
        match_type = "UNMATCHED"
        
        p_amt = p["amount"]
        p_narr = p["narration"].strip()
        p_narr_lower = p_narr.lower()
        p_bank_lower = p["bank_name"].strip().lower()
        
        # 1. Exact match on narration and amount
        for tx in account_transactions:
            if tx["id"] in used_tx_ids: continue
            if abs(tx["amount"] - p_amt) < 0.01:
                tx_narr = tx["narration"].strip()
                if tx_narr and (p_narr == tx_narr or p_narr_lower in tx_narr.lower() or tx_narr.lower() in p_narr_lower):
                    matched_tx = tx
                    match_type = "EXACT_DESCRIPTION"
                    break
                    
        # 2. Match on UTR / Cheque number if extracted
        if not matched_tx:
            p_cheque = extract_cheque_no(p_narr)
            if p_cheque and len(p_cheque) >= 6:
                for tx in account_transactions:
                    if tx["id"] in used_tx_ids: continue
                    tx_cheque = extract_cheque_no(tx["narration"], tx["cheque_ref"])
                    if tx_cheque == p_cheque and abs(tx["amount"] - p_amt) < 0.01:
                        matched_tx = tx
                        match_type = "MATCH_UTR"
                        break
                        
        # 3. Match on Amount + Bank
        if not matched_tx and p_bank_lower:
            candidates = [
                tx for tx in account_transactions 
                if tx["id"] not in used_tx_ids 
                and abs(tx["amount"] - p_amt) < 0.01 
                and (p_bank_lower in tx["bank"].lower() or tx["bank"].lower() in p_bank_lower)
            ]
            if len(candidates) == 1:
                matched_tx = candidates[0]
                match_type = "AMOUNT_BANK_MATCH"
                
        # 4. Match on Amount only (if unique)
        if not matched_tx:
            candidates = [
                tx for tx in account_transactions 
                if tx["id"] not in used_tx_ids 
                and abs(tx["amount"] - p_amt) < 0.01
            ]
            if len(candidates) == 1:
                matched_tx = candidates[0]
                match_type = "UNIQUE_AMOUNT_MATCH"
                
        # Derive final values
        ref_no = ""
        bank_final = p["bank_name"]
        final_narr = p["narration"]
        raw_cheque = ""
        
        if matched_tx:
            used_tx_ids.add(matched_tx["id"])
            ref_no = matched_tx["ref"]
            if not bank_final:
                bank_final = matched_tx["bank"]
            if not final_narr:
                final_narr = matched_tx["narration"]
            raw_cheque = matched_tx["cheque_ref"]
        else:
            # Generate temporary fallback reference if completely unmatched
            ref_no = f"AUTO-{p_idx+1:04d}"
            
        cheque_no = extract_cheque_no(final_narr, raw_cheque)
        mode_final = detect_mode(p["mode"], final_narr)
        
        # Portfolio mapping
        port_norm = p["portfolio"].strip().upper()
        port_info = portfolio_map.get(port_norm, {})
        port_source = port_info.get("source", port_norm)
        prod_type = port_info.get("product", "PL" if "PL" in port_norm else ("CC" if "CC" in port_norm else "PL"))
        
        # Agency mapping
        mgr_norm = p["manager"].strip().upper()
        agency = manager_agency_map.get(mgr_norm, default_agency)
        if p["officer_name"].strip().upper() == "KP" or mgr_norm == "KP":
            agency = "KP"
            
        handsoff_results.append({
            "PORTFOLIO": p["portfolio"],
            "PORTFOLIO TYPE/SOURCE": port_source,
            "PRODUCT TYPE": prod_type,
            "Loan ID": p["customer_id"],
            "LOAN NO": p["customer_id"],
            "Customer Name": p["customer_name"],
            "Mode of Pay": mode_final,
            "Date of Pay": p["date_of_pay"],
            "Reference No": ref_no,
            "Amount": p["amount"],
            "CHEQUE.no": cheque_no,
            "TRANSACTION DESCRIPTION": final_narr,
            "HAND OFF DATE": handoff_date,
            "AGENCY NAME": agency,
            "CALLER NAME": p["officer_name"],
            "AGENCY MANAGER": p["manager"],
            "BANK": bank_final,
            "match_status": match_type,
            "source_file": p.get("source_file", "")
        })
        
    return handsoff_results, len(used_tx_ids), len(paid_rows)

def write_handsoff_excel(results, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    
    headers = [
        'PORTFOLIO', 'PORTFOLIO TYPE/SOURCE', 'PRODUCT TYPE', 'Loan ID', 'LOAN NO',
        'Customer Name', 'Mode of Pay', 'Date of Pay', 'Reference No', 'Amount',
        'CHEQUE.no', 'TRANSACTION DESCRIPTION', 'HAND OFF DATE', 'AGENCY NAME',
        'CALLER NAME', 'AGENCY MANAGER', 'BANK', 'DAY', 'CONCATE', 'REMARKS'
    ]
    
    col_widths = {
        'A': 15.0, 'B': 25.0, 'C': 15.0, 'D': 18.0, 'E': 16.0,
        'F': 28.0, 'G': 12.0, 'H': 14.0, 'I': 16.0, 'J': 12.0,
        'K': 18.0, 'L': 65.0, 'M': 15.0, 'N': 15.0, 'O': 22.0,
        'P': 20.0, 'Q': 10.0, 'R': 12.0, 'S': 24.0, 'T': 15.0
    }
    
    font_header = Font(name="Calibri", size=11, bold=True, color="000000")
    fill_header = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    font_data = Font(name="Calibri", size=11, color="000000")
    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    # Write header
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        cell.border = border_thin
        
    for r_idx, item in enumerate(results, start=2):
        row_vals = [
            item.get('PORTFOLIO', ''),
            item.get('PORTFOLIO TYPE/SOURCE', ''),
            item.get('PRODUCT TYPE', ''),
            item.get('Loan ID', ''),
            item.get('LOAN NO', ''),
            item.get('Customer Name', ''),
            item.get('Mode of Pay', ''),
            item.get('Date of Pay', ''),
            item.get('Reference No', ''),
            item.get('Amount', 0.0),
            item.get('CHEQUE.no', ''),
            item.get('TRANSACTION DESCRIPTION', ''),
            item.get('HAND OFF DATE', ''),
            item.get('AGENCY NAME', ''),
            item.get('CALLER NAME', ''),
            item.get('AGENCY MANAGER', ''),
            item.get('BANK', ''),
            f'=TEXT(H{r_idx},"DDDDDD")',  # DAY formula
            f'=N{r_idx}&O{r_idx}',        # CONCATE formula
            item.get('REMARKS', '')
        ]
        
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_data
            cell.border = border_thin
            
            # Alignments
            if c_idx in (7, 8, 9, 13, 14, 17, 18):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 10: # Amount
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0.00"
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
    # Apply Column Widths
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width
        
    ws.freeze_panes = "A2"
    wb.save(output_path)
    wb.close()
    return output_path

def write_paid_excel(paid_rows, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Paid File"
    
    headers = [
        'CUSTOMER ID ', 'CUSTOMER  NAME', 'PORTFOLIO', 'Date of Pay', 'Amount',
        'MODE', 'TRANSACTION DESCRIPTION', 'BANK NAME', 'OFFICER NAME', 'MANAGER'
    ]
    
    font_header = Font(name="Calibri", size=11, bold=True, color="000000")
    fill_header = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    font_data = Font(name="Calibri", size=11, color="000000")
    
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    for r_idx, item in enumerate(paid_rows, start=2):
        row_vals = [
            item.get('customer_id', ''),
            item.get('customer_name', ''),
            item.get('portfolio', ''),
            item.get('date_of_pay', ''),
            item.get('amount', 0.0),
            item.get('mode', ''),
            item.get('narration', ''),
            item.get('bank_name', ''),
            item.get('officer_name', ''),
            item.get('manager', '')
        ]
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_data
            if c_idx in (4, 6, 8):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 5:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0.00"
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
    ws.column_dimensions['A'].width = 18.0
    ws.column_dimensions['B'].width = 25.0
    ws.column_dimensions['C'].width = 15.0
    ws.column_dimensions['D'].width = 14.0
    ws.column_dimensions['E'].width = 14.0
    ws.column_dimensions['F'].width = 12.0
    ws.column_dimensions['G'].width = 50.0
    ws.column_dimensions['H'].width = 15.0
    ws.column_dimensions['I'].width = 20.0
    ws.column_dimensions['J'].width = 20.0
    
    wb.save(output_path)
    wb.close()
    return output_path
